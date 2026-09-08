from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from bot.repositories import submissions as submissions_repo

HEADERS = (
    "id",
    "telegram_id",
    "username",
    "имя",
    "день",
    "дата",
    "попытка",
    "длительность",
    "отправлено",
    "статус",
    "причина отказа",
    "решено",
    "флаги",
    "file_unique_id",
    "пост в канале",
    "судья (user.id)",
)


def _cell(value: object, tz: ZoneInfo) -> object:
    if isinstance(value, datetime):
        return value.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if hasattr(value, "value"):  # StrEnum
        return value.value
    return value


async def build_xlsx(session: AsyncSession, challenge_id: int, tz: ZoneInfo) -> bytes | None:
    rows = await submissions_repo.export_rows(session, challenge_id)
    if not rows:
        return None

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сдачи"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append([_cell(v, tz) for v in row])

    sheet.freeze_panes = "A2"
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) + 4)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
